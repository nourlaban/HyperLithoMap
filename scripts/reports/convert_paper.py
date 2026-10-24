#!/usr/bin/env python3
"""Convert RESEARCH_PAPER_DRAFT.md to a formatted DOCX file."""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── helpers ──────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_run_with_style(para, text, bold=False, italic=False, code=False,
                       font_size=None, color=None):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    if code:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
    if font_size:
        run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = color
    return run


def parse_inline(para, text):
    """Parse bold/italic/inline-code markup and add styled runs."""
    # Combined regex: **bold**, *italic*, `code`
    pattern = re.compile(r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
    last = 0
    for m in pattern.finditer(text):
        # plain text before match
        if m.start() > last:
            para.add_run(text[last:m.start()])
        if m.group(0).startswith('**'):
            add_run_with_style(para, m.group(2), bold=True)
        elif m.group(0).startswith('*'):
            add_run_with_style(para, m.group(3), italic=True)
        elif m.group(0).startswith('`'):
            add_run_with_style(para, m.group(4), code=True)
        last = m.end()
    # remaining plain text
    if last < len(text):
        para.add_run(text[last:])


def add_heading(doc, text, level):
    """Add a heading at the given level (1-4)."""
    style_map = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3', 4: 'Heading 4'}
    style = style_map.get(level, 'Heading 3')
    try:
        h = doc.add_heading(text, level=level)
    except Exception:
        h = doc.add_paragraph(text, style=style)
    return h


def add_table(doc, rows):
    """
    rows: list of lists of cell strings.
    rows[0] = header row.
    """
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Table Grid'

    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            if ci >= cols:
                break
            cell = table.cell(ri, ci)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            # style header row
            if ri == 0:
                set_cell_bg(cell, 'D6E4F0')
                p = cell.paragraphs[0]
                p.clear()
                run = p.add_run(cell_text.strip())
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
            else:
                p = cell.paragraphs[0]
                p.clear()
                parse_inline(p, cell_text.strip())
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()   # spacing after table


def parse_table_row(line):
    """Split a markdown table row into cell strings."""
    line = line.strip().strip('|')
    cells = [c.strip() for c in line.split('|')]
    return cells


def is_separator_row(line):
    """True if line is a markdown table separator like |---|---|"""
    stripped = line.strip().strip('|')
    return bool(re.fullmatch(r'[\s\-|:]+', stripped))


# ── main converter ────────────────────────────────────────────────────────────

def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # ── document-wide styles ──────────────────────────────────────────────────
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Heading colours
    heading_colors = {
        1: RGBColor(0x1A, 0x3A, 0x5C),   # dark navy
        2: RGBColor(0x1F, 0x61, 0x9E),   # medium blue
        3: RGBColor(0x2E, 0x86, 0xAB),   # teal-blue
        4: RGBColor(0x44, 0x88, 0x6E),   # green
    }
    for level, color in heading_colors.items():
        hname = f'Heading {level}'
        if hname in doc.styles:
            doc.styles[hname].font.color.rgb = color
            doc.styles[hname].font.bold = True
            sizes = {1: 18, 2: 15, 3: 13, 4: 11}
            doc.styles[hname].font.size = Pt(sizes[level])

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)

    # ── state machine ─────────────────────────────────────────────────────────
    i = 0
    in_code_block = False
    code_lines = []
    table_rows = []
    in_table = False
    bullet_buffer = []

    def flush_bullets():
        if bullet_buffer:
            for item in bullet_buffer:
                p = doc.add_paragraph(style='List Bullet')
                parse_inline(p, item)
                p.paragraph_format.space_after = Pt(2)
            bullet_buffer.clear()
            doc.add_paragraph()

    def flush_table():
        if table_rows:
            add_table(doc, table_rows)
            table_rows.clear()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip('\n')
        stripped = line.strip()

        # ── code block ────────────────────────────────────────────────────────
        if stripped.startswith('```'):
            if not in_code_block:
                flush_bullets()
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                # add code block as styled paragraph(s)
                for cl in code_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(1)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    run = p.add_run(cl)
                    run.font.name = 'Courier New'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(0x3D, 0x40, 0x45)
                doc.add_paragraph()
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # ── horizontal rule ───────────────────────────────────────────────────
        if re.fullmatch(r'-{3,}', stripped):
            flush_bullets()
            flush_table()
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'AAAAAA')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # ── headings ──────────────────────────────────────────────────────────
        hm = re.match(r'^(#{1,4})\s+(.*)', line)
        if hm:
            flush_bullets()
            flush_table()
            level = len(hm.group(1))
            text = hm.group(2).strip()
            add_heading(doc, text, level)
            i += 1
            continue

        # ── blockquote (> text) — used for figure captions ────────────────────
        if stripped.startswith('> '):
            flush_bullets()
            flush_table()
            content = stripped[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(content)
            run.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ── markdown table ────────────────────────────────────────────────────
        if stripped.startswith('|'):
            if is_separator_row(stripped):
                i += 1
                continue
            flush_bullets()
            table_rows.append(parse_table_row(stripped))
            i += 1
            continue
        else:
            if table_rows:
                flush_table()

        # ── bullet list ───────────────────────────────────────────────────────
        bm = re.match(r'^(\s*[-*])\s+(.*)', line)
        if bm:
            flush_table()
            bullet_buffer.append(bm.group(2))
            i += 1
            continue
        else:
            flush_bullets()

        # ── numbered list ─────────────────────────────────────────────────────
        nm = re.match(r'^\s*\d+\.\s+(.*)', line)
        if nm:
            flush_table()
            p = doc.add_paragraph(style='List Number')
            parse_inline(p, nm.group(1))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # ── blank line ────────────────────────────────────────────────────────
        if stripped == '':
            flush_table()
            flush_bullets()
            i += 1
            continue

        # ── paper title (first H1 equiv — bold italic large) ──────────────────
        # Already handled by heading block above.

        # ── metadata lines (Author, Affiliation, etc.) ────────────────────────
        if stripped.startswith('**Author**') or stripped.startswith('**Affiliation**') \
                or stripped.startswith('**Keywords**') or stripped.startswith('**Suggested Journal**'):
            p = doc.add_paragraph()
            parse_inline(p, stripped)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # ── math block ($$...$$) — render as styled paragraph ─────────────────
        if stripped.startswith('$$'):
            # collect until closing $$
            math_lines = []
            if stripped.count('$$') >= 2 and stripped != '$$':
                # single-line
                math_lines.append(stripped.replace('$$', '').strip())
            else:
                i += 1
                while i < len(lines):
                    ml = lines[i].rstrip('\n').strip()
                    if ml == '$$':
                        break
                    math_lines.append(ml)
                    i += 1
            math_text = ' '.join(math_lines)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(math_text)
            run.font.name = 'Cambria Math'
            run.font.size = Pt(10.5)
            run.italic = True
            run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
            i += 1
            continue

        # ── regular paragraph ─────────────────────────────────────────────────
        # Check if it's a "Note:" line
        if stripped.lower().startswith('*note:') or stripped.lower().startswith('> **note'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            run = p.add_run(stripped.strip('*').strip())
            run.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            parse_inline(p, stripped)

        i += 1

    # flush any trailing state
    flush_bullets()
    flush_table()

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(description='Convert a Markdown paper draft to a formatted DOCX.')
    ap.add_argument('markdown', help='Input Markdown file, e.g. RESEARCH_PAPER_DRAFT.md')
    ap.add_argument('docx', help='Output .docx path')
    a = ap.parse_args()
    convert(a.markdown, a.docx)
